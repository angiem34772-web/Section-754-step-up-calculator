"""Write the step-up calculation to a formatted Excel workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Results

MONEY = '#,##0;(#,##0);"-"'
PCT = "0.00%"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")

TYPE_LABELS = {
    "cash": "Cash",
    "receivable": "Receivable",
    "inventory": "Inventory",
    "land": "Land",
    "residential_building": "Residential rental building",
    "nonresidential_building": "Nonresidential building",
    "personal_property": "Section 1245 personal property",
    "capital_asset": "Capital asset",
    "section_197": "Section 197 intangible",
}

RECOVERY_LABELS = {
    "SL-MM": "Straight line, mid-month",
    "200DB-HY": "MACRS declining balance, half-year",
    "SL-197": "15-year ratable amortization",
}


def export(results: Results, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary(wb.active, results)
    _calculation(wb.create_sheet("743(b) Calculation"), results)
    _hypothetical_sale(wb.create_sheet("Hypothetical Sale"), results)
    _allocation(wb.create_sheet("755 Allocation"), results)
    _depreciation(wb.create_sheet("Depreciation Schedule"), results)
    if results.seller:
        _seller(wb.create_sheet("Seller Gain"), results)
    _inputs(wb.create_sheet("Inputs"), results)
    wb.save(path)
    return path


def _header(ws, headers, row=1, height=45):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row, col, text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def _table(ws, headers, rows, money_from: int, widths=None, total_row: bool = False):
    _header(ws, headers)
    for row in rows:
        ws.append([round(v, 2) if isinstance(v, float) else v for v in row])
    for row in ws.iter_rows(min_row=2, min_col=money_from):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
    if total_row:
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = Font(bold=True), TOTAL_FILL
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(i, max(12, min(len(str(header)) + 2, 18)))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"


def _summary(ws, r: Results):
    s = r.sale
    ws.title = "Summary"
    ws["A1"] = "Section 754 Step-Up Calculator"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Sample data only. The partnership, partners, and amounts are fictional."
    ws["A2"].font = Font(italic=True, color="7F7F7F")

    years = r.years
    ten = [y for y in years if y < s.sale_date.year + 10]
    lines = [
        ("The sale", None),
        ("Partnership", s.partnership),
        ("Seller", s.seller),
        ("Buyer", s.buyer),
        ("Sale date", s.sale_date.isoformat()),
        ("Interest sold", s.interest_sold),
        ("Cash price", s.cash_price),
        ("754 election in effect", "Yes" if s.election_754 else "No"),
        ("", None),
        ("The result", None),
        ("Buyer's outside basis", r.outside_basis),
        ("Buyer's share of inside basis", r.inside_basis_share),
        ("743(b) adjustment", r.adjustment),
        ("  to ordinary income property", r.ordinary_class),
        ("  to capital gain property", r.capital_class),
        ("Adjustment applies?", "Yes" if r.adjustment_applies else "No (no 754 election)"),
        ("", None),
        ("Why it matters to the buyer", None),
        (f"Extra deductions in {s.sale_date.year}", r.deduction(s.sale_date.year)),
        (f"Extra deductions, first 10 years", sum(r.deduction(y) for y in ten)),
        (f"Extra deductions, full recovery ({years[0]}-{years[-1]})" if years else "Extra deductions",
         sum(r.deduction(y) for y in years)),
        (f"Illustrative tax saved, first 10 years at {s.tax_rate:.0%}",
         sum(r.deduction(y) for y in ten) * s.tax_rate),
    ]
    row = 4
    for label, value in lines:
        cell = ws.cell(row, 1, label)
        if value is None and label:
            cell.font = Font(bold=True, size=12)
        elif label:
            cell.font = Font(bold=True)
            out = ws.cell(row, 2, round(value, 2) if isinstance(value, float) else value)
            if label == "Interest sold":
                out.number_format = PCT
            elif isinstance(value, (int, float)):
                out.number_format = MONEY
            if label == "743(b) adjustment":
                out.fill = cell.fill = KEY_FILL
            out.alignment = Alignment(horizontal="right")
        row += 1

    row += 1
    ws.cell(row, 1, "Flags").font = Font(bold=True, size=12)
    for flag in r.flags:
        row += 1
        ws.cell(row, 1, flag.code)
        ws.cell(row, 2, round(flag.amount, 2) if flag.amount is not None else None).number_format = MONEY
        ws.cell(row, 3, flag.message).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30

    row += 2
    ws.cell(row, 1, "How to read this workbook").font = Font(bold=True, size=12)
    guide = [
        ("743(b) Calculation", "The adjustment, step by step, with the regulation behind each line."),
        ("Hypothetical Sale", "Each asset 'sold' at fair market value, and the buyer's share of the gain or loss."),
        ("755 Allocation", "How the adjustment is split between ordinary and capital assets, then by asset."),
        ("Depreciation Schedule", "The buyer's extra deductions from the step-up, year by year."),
        ("Seller Gain", "The seller's side: total gain and its character (ordinary, 25%, 20%)."),
        ("Inputs", "The balance sheet and sale terms the calculation started from."),
    ]
    for sheet, text in guide:
        row += 1
        ws.cell(row, 1, sheet).font = Font(bold=True)
        ws.cell(row, 3, text)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 95


def _calculation(ws, r: Results):
    s = r.sale
    pct = s.interest_sold
    rows = [
        ("A", "Cash paid by the buyer", s.cash_price, "IRC 742, 1012"),
        ("B", "Buyer's share of partnership liabilities", s.debt_share, "IRC 752(a), 752(d)"),
        ("C", "Buyer's outside basis (A + B)", r.outside_basis, "IRC 742"),
        ("", "", None, ""),
        ("D", "Partnership gross value: A / interest sold + total liabilities", r.gross_value, "Reg. 1.755-1(a)(4)"),
        ("E", "Fair market value of the listed assets", r.gross_value - r.residual_goodwill, ""),
        ("F", "Residual value assigned to goodwill (D - E)", r.residual_goodwill, "Reg. 1.755-1(a)(5)"),
        ("", "", None, ""),
        ("G", "Cash the buyer would receive if the partnership sold everything and liquidated",
         r.cash_on_liquidation, "Reg. 1.743-1(d)(2)(i)"),
        ("H", "Buyer's share of gain/(loss) from that hypothetical sale", r.share_of_gain, "Reg. 1.743-1(d)(2)"),
        ("I", "Previously taxed capital (G - H)", r.previously_taxed_capital, "Reg. 1.743-1(d)(2)"),
        ("J", "Buyer's share of partnership liabilities (B)", s.debt_share, "Reg. 1.743-1(d)(1)"),
        ("K", "Buyer's share of inside basis (I + J)", r.inside_basis_share, "Reg. 1.743-1(d)(1)"),
        ("", "", None, ""),
        ("L", "743(b) adjustment (C - K)", r.adjustment, "IRC 743(b), Reg. 1.743-1(b)"),
        ("", "", None, ""),
        ("M", "Allocated to ordinary income property", r.ordinary_class, "Reg. 1.755-1(b)(2)(i)"),
        ("N", "Allocated to capital gain property (L - M)", r.capital_class, "Reg. 1.755-1(b)(2)(i)"),
        ("", "", None, ""),
        ("", f"Check: buyer's share of the tax capital account ({pct:.0%} x (basis - liabilities))",
         pct * (sum(a.tax_basis for a in r.assets) - s.partnership_liabilities), "Equals I with pro rata allocations"),
    ]
    _table(ws, ["Line", "Step", "Amount", "Authority"], rows, money_from=3, widths={1: 6, 2: 80, 3: 16, 4: 34})
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "L":
            for cell in row:
                cell.font, cell.fill = Font(bold=True), KEY_FILL
    ws.freeze_panes = None


def _hypothetical_sale(ws, r: Results):
    pct = r.sale.interest_sold
    lines_by_asset = {}
    for line in r.lines:
        lines_by_asset.setdefault(line.asset_id, []).append(line)
    rows = []
    for a in r.assets:
        ls = lines_by_asset.get(a.asset_id, [])
        cls = " / ".join(sorted({l.asset_class.title() for l in ls})) or "Not adjusted (money)"
        rows.append([a.asset_id, a.description, TYPE_LABELS[a.asset_type], a.tax_basis, a.fair_market_value,
                     a.gain, a.prior_depreciation, a.recapture, pct * a.recapture, pct * (a.gain - a.recapture),
                     pct * a.gain, cls])
    rows.append(["Total", "", "", sum(a.tax_basis for a in r.assets), sum(a.fair_market_value for a in r.assets),
                 sum(a.gain for a in r.assets), sum(a.prior_depreciation for a in r.assets),
                 sum(a.recapture for a in r.assets), sum(pct * a.recapture for a in r.assets),
                 sum(pct * (a.gain - a.recapture) for a in r.assets), r.share_of_gain, ""])
    headers = ["Asset", "Description", "Type", "Partnership tax basis", "Fair market value", "Gain / (loss)",
               "Depreciation taken to date", "1245 recapture (ordinary)", f"Buyer's {pct:.0%} share: ordinary",
               f"Buyer's {pct:.0%} share: capital", f"Buyer's {pct:.0%} share: total", "755 class"]
    _table(ws, headers, rows, money_from=4, widths={2: 34, 3: 28, 12: 20}, total_row=True)


def _allocation(ws, r: Results):
    pct = r.sale.interest_sold
    rows = []
    for cls, total in (("ordinary", r.ordinary_class), ("capital", r.capital_class)):
        items = [l for l in r.lines if l.asset_class == cls]
        for l in items:
            rows.append([l.asset_id, l.description, cls.title(), l.basis, l.fmv, l.share_of_gain,
                         l.adjustment - l.share_of_gain, l.adjustment, pct * l.basis, pct * l.basis + l.adjustment])
        rows.append([f"{cls.title()} class total", "", "", sum(l.basis for l in items), sum(l.fmv for l in items),
                     sum(l.share_of_gain for l in items), total - sum(l.share_of_gain for l in items), total,
                     sum(pct * l.basis for l in items), sum(pct * l.basis + l.adjustment for l in items)])
    rows.append(["743(b) adjustment", "", "", None, None, None, None, r.adjustment, None, None])
    headers = ["Asset", "Item", "755 class", "Partnership basis in item", "Fair market value",
               "Buyer's share of gain / (loss)", "Pro rata true-up (Reg. 1.755-1(b)(3))", "743(b) adjustment",
               "Buyer's share of common basis", "Buyer's basis after step-up"]
    _table(ws, headers, rows, money_from=4, widths={1: 22, 2: 40})
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).endswith("total") or row[0].value == "743(b) adjustment":
            for cell in row:
                cell.font, cell.fill = Font(bold=True), TOTAL_FILL


def _depreciation(ws, r: Results):
    names = {a.asset_id: a for a in r.assets}
    ids = list(r.depreciation)
    rate = r.sale.tax_rate
    _header(ws, ["Tax year"] + [f"{names[i].description} ({i})" for i in ids]
            + ["Total extra deduction", "Cumulative", f"Illustrative tax saved at {rate:.0%}"])
    ws.append(["Step-up"] + [round(r.by_asset[i], 2) for i in ids] + [round(sum(r.by_asset[i] for i in ids), 2)])
    ws.append(["Method"] + [RECOVERY_LABELS[names[i].recovery] if r.by_asset[i] > 0 else "Remaining period"
                            for i in ids])
    ws.append(["Recovery period (years)"] + [names[i].years if r.by_asset[i] > 0 else names[i].remaining_years
                                             for i in ids])
    for row in ws.iter_rows(min_row=2, max_row=4):
        for cell in row:
            cell.font = Font(italic=True)
    cumulative = 0.0
    for year in r.years:
        total = r.deduction(year)
        cumulative += total
        ws.append([year] + [round(r.depreciation[i].get(year, 0.0), 2) for i in ids]
                  + [round(total, 2), round(cumulative, 2), round(total * rate, 2)])
    ws.append(["Total"] + [round(sum(r.depreciation[i].values()), 2) for i in ids]
              + [round(cumulative, 2), None, round(cumulative * rate, 2)])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = Font(bold=True), TOTAL_FILL
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = MONEY
    for row in ws.iter_rows(min_row=4, max_row=4, min_col=2):
        for cell in row:
            cell.number_format = "0.0"
    ws.column_dimensions["A"].width = 24
    for col in range(2, len(ids) + 5):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.freeze_panes = "B5"


def _seller(ws, r: Results):
    s = r.seller
    rows = [
        ("Cash received", s["cash_received"], "IRC 1001"),
        ("Relief from share of partnership liabilities", s["debt_relief"], "IRC 752(d)"),
        ("Amount realized", s["amount_realized"], ""),
        ("Seller's outside basis", -s["outside_basis"], "From the seller's basis records"),
        ("Total gain on sale", s["total_gain"], "IRC 741"),
        ("", None, ""),
        ("Ordinary income from hot assets", s["ordinary_751"], "IRC 751(a), Reg. 1.751-1(a)(2)"),
        ("Capital gain", s["capital_gain"], "IRC 741"),
        ("  of which unrecaptured 1250 gain (25% rate)", s["unrecaptured_1250"], "IRC 1(h)(6), Reg. 1.1(h)-1"),
        ("  of which other long-term capital gain", s["other_long_term"], "Assumes interest held over 1 year"),
    ]
    _table(ws, ["Seller: " + r.sale.seller, "Amount", "Authority"], rows, money_from=2,
           widths={1: 50, 2: 16, 3: 40})
    ws.freeze_panes = None


def _inputs(ws, r: Results):
    s = r.sale
    rows = [[a.asset_id, a.description, a.asset_type, a.tax_basis, a.fair_market_value, a.prior_depreciation,
             a.recovery_years, a.remaining_years] for a in r.assets]
    headers = ["asset_id", "description", "asset_type", "tax_basis", "fair_market_value", "prior_depreciation",
               "recovery_years", "remaining_years"]
    _table(ws, headers, rows, money_from=4, widths={2: 40, 3: 24})
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = "0.0"
    start = ws.max_row + 3
    ws.cell(start - 1, 1, "Sale terms").font = Font(bold=True, size=12)
    terms = [("partnership", s.partnership), ("seller", s.seller), ("buyer", s.buyer),
             ("sale_date", s.sale_date.isoformat()), ("interest_sold", s.interest_sold),
             ("cash_price", s.cash_price), ("partnership_liabilities", s.partnership_liabilities),
             ("liability_share", s.liability_share), ("seller_outside_basis", s.seller_outside_basis),
             ("election_754", "yes" if s.election_754 else "no"), ("bonus_pct", s.bonus_pct),
             ("tax_rate", s.tax_rate)]
    for i, (k, v) in enumerate(terms):
        ws.cell(start + i, 1, k).font = Font(bold=True)
        cell = ws.cell(start + i, 2, v)
        if k in ("cash_price", "partnership_liabilities", "seller_outside_basis"):
            cell.number_format = MONEY
    if r.residual_goodwill and any(a.asset_id == "GW" for a in r.assets):
        ws.cell(start + len(terms) + 1, 1, "Goodwill row 'GW' is computed from the price (residual method), "
                "not entered.").font = Font(italic=True, color="7F7F7F")
